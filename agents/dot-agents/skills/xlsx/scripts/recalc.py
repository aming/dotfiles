"""
Excel Formula Recalculation Script
Recalculates all formulas in an Excel file using LibreOffice
"""

import json
import tempfile
import zipfile
import defusedxml.ElementTree as ET
import uuid
import subprocess
import sys
from pathlib import Path

from office.soffice import get_soffice_env

from openpyxl import load_workbook

RECALCULATE_MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="{module}" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      Dim marker As Integer
      marker = FreeFile
      Open "{marker}" For Output As #marker
      Print #marker, "recalculated-and-saved"
      Close #marker
      ThisComponent.close(True)
    End Sub
</script:module>"""


def setup_libreoffice_macro(profile, module, marker, timeout):
    result = subprocess.run(
        ["soffice", "--headless", f"-env:UserInstallation={profile.as_uri()}", "--terminate_after_init"],
        capture_output=True, timeout=timeout, env=get_soffice_env(),
    )
    if result.returncode != 0:
        raise RuntimeError("LibreOffice profile initialization failed")
    macro_dir = profile / "user/basic/Standard"
    macro_dir.mkdir(parents=True, exist_ok=True)
    (macro_dir / f"{module}.xba").write_text(RECALCULATE_MACRO.format(module=module, marker=marker))
    # Register the unique module in the private Standard library.
    (macro_dir / "script.xlb").write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<library:library xmlns:library="http://openoffice.org/2000/library" library:name="Standard" library:readonly="false" library:passwordprotected="false">'
        f'<library:element library:name="{module}"/></library:library>'
    )


def recalc(filename, timeout=30):
    if not Path(filename).exists():
        return {"error": f"File {filename} does not exist"}

    if timeout <= 0:
        return {"error": "Timeout must be positive"}

    abs_path = str(Path(filename).absolute())

    try:
        with tempfile.TemporaryDirectory(prefix="xlsx-recalc-") as temp_dir:
            profile = Path(temp_dir) / "profile"
            marker = Path(temp_dir) / "completed"
            module = "Recalc_" + uuid.uuid4().hex
            setup_libreoffice_macro(profile, module, marker, timeout)
            cmd = [
                "soffice", "--headless", f"-env:UserInstallation={profile.as_uri()}",
                "--norestore",
                f"vnd.sun.star.script:Standard.{module}.RecalculateAndSave?language=Basic&location=application",
                abs_path,
            ]
            process = subprocess.run(cmd, capture_output=True, text=True,
                                     timeout=timeout, env=get_soffice_env())
            if process.returncode != 0:
                return {"error": f"LibreOffice failed ({process.returncode}): {process.stderr}"}
            if not marker.is_file() or marker.read_text().strip() != "recalculated-and-saved":
                return {"error": "Recalculation completion was not confirmed by the macro"}
    except subprocess.TimeoutExpired:
        return {"error": f"LibreOffice timed out after {timeout} seconds; recalculation unverified"}
    except Exception as exc:
        return {"error": str(exc)}

    try:
        missing_caches = []
        sheet_namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        with zipfile.ZipFile(filename) as archive:
            for part in archive.namelist():
                if not part.startswith("xl/") or not part.endswith(".xml"):
                    continue
                root = ET.fromstring(archive.read(part))
                if root.tag != sheet_namespace + "worksheet":
                    continue
                for cell in root.iter(sheet_namespace + "c"):
                    if cell.find(sheet_namespace + "f") is None:
                        continue
                    value = cell.find(sheet_namespace + "v")
                    # A cached empty string is valid; an absent/empty numeric cache is not.
                    if value is None or (value.text is None and cell.get("t") != "str"):
                        missing_caches.append(f"{part}!{cell.get('r')}")
        if missing_caches:
            return {"error": "Formula cache coverage is incomplete after recalculation",
                    "missing_formula_caches": missing_caches[:20],
                    "total_missing_formula_caches": len(missing_caches)}

        wb = load_workbook(filename, data_only=True, keep_vba=Path(filename).suffix.lower() == ".xlsm")

        excel_errors = [
            "#VALUE!",
            "#DIV/0!",
            "#REF!",
            "#NAME?",
            "#NULL!",
            "#NUM!",
            "#N/A",
        ]
        error_details = {err: [] for err in excel_errors}
        total_errors = 0

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "e":
                        error_details.setdefault(cell.value, []).append(f"{sheet_name}!{cell.coordinate}")
                        total_errors += 1

        wb.close()

        result = {
            "status": "success" if total_errors == 0 else "errors_found",
            "total_errors": total_errors,
            "error_summary": {},
        }

        for err_type, locations in error_details.items():
            if locations:
                result["error_summary"][err_type] = {
                    "count": len(locations),
                    "locations": locations[:20],  
                }

        wb_formulas = load_workbook(filename, data_only=False, keep_vba=Path(filename).suffix.lower() == ".xlsm")
        formula_count = 0
        for sheet_name in wb_formulas.sheetnames:
            ws = wb_formulas[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula_count += 1
        wb_formulas.close()

        result["total_formulas"] = formula_count

        return result

    except Exception as e:
        return {"error": str(e)}


def main():
    if len(sys.argv) < 2:
        print("Usage: python recalc.py <excel_file> [timeout_seconds]")
        print("\nRecalculates all formulas in an Excel file using LibreOffice")
        print("\nReturns JSON with error details:")
        print("  - status: 'success' or 'errors_found'")
        print("  - total_errors: Total number of Excel errors found")
        print("  - total_formulas: Number of formulas in the file")
        print("  - error_summary: Breakdown by error type with locations")
        print("    - #VALUE!, #DIV/0!, #REF!, #NAME?, #NULL!, #NUM!, #N/A")
        sys.exit(1)

    filename = sys.argv[1]
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 30

    result = recalc(filename, timeout)
    print(json.dumps(result, indent=2))
    sys.exit(1 if "error" in result or result.get("total_errors", 0) else 0)


if __name__ == "__main__":
    main()
